import io
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from report.layout import get_styles, get_table_style
from report.templates import add_header_footer


def formatear_importe(amount):
    return f"{amount:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def construir_filas_reporte(records):
    running_total = 0.0
    rows_by_id = {}

    for record in sorted(records, key=lambda item: item.date):
        running_total += record.amount
        rows_by_id[record.id] = {
            'date': record.date.strftime('%d/%m/%Y'),
            'concept': record.concept,
            'description': record.description or "",
            'amount': float(record.amount),
            'running_total': running_total,
        }

    ordered_rows = [rows_by_id[record.id] for record in records]
    return ordered_rows, running_total


def generar_pdf(datos_lista, filas_reporte, saldo_total, nombre):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    elementos = []
    styles = get_styles()

    # Título
    elementos.append(Paragraph(f'Caja {nombre}', styles["CenterTitle"]))
    elementos.append(Spacer(1, 12))

    # Tabla con saldo como primera fila
    encabezado = ["Fecha", "Concepto", "Descripción", "Cantidad", "Total acumulado"]
    tabla_datos = [encabezado]

    # Fila de saldo destacada
    saldo_formateado = formatear_importe(saldo_total)
    saldo_color = colors.red if saldo_total < 0 else colors.green
    tabla_datos.append([
        Paragraph("<b>Saldo:</b>", styles["NormalText"]),
        "", "", "",
        Paragraph(f'<font color="{saldo_color}"><b>{saldo_formateado}</b></font>', styles["NormalText"])
    ])

    # Resto de registros
    for fila in filas_reporte:
        cantidad_num = fila['amount']
        running_total = fila['running_total']
        cantidad_color = colors.red if cantidad_num < 0 else colors.black
        acumulado_color = colors.red if running_total < 0 else colors.green
        cantidad_formateada = formatear_importe(cantidad_num)
        acumulado_formateado = formatear_importe(running_total)

        tabla_datos.append([
            fila['date'],
            fila['concept'],
            fila['description'],
            Paragraph(f'<font color="{cantidad_color}">{cantidad_formateada}</font>', styles["NormalText"]),
            Paragraph(f'<font color="{acumulado_color}">{acumulado_formateado}</font>', styles["NormalText"])
        ])

    tabla = Table(tabla_datos, colWidths=[2.7*cm, 4.4*cm, 5.2*cm, 2.8*cm, 2.9*cm])
    tabla.setStyle(TableStyle(get_table_style()))
    elementos.append(tabla)

    doc.build(elementos, onFirstPage=add_header_footer, onLaterPages=add_header_footer)
    buffer.seek(0)
    return buffer


def _build_sheet_name(person_name, used_names):
    invalid_chars = {'\\', '/', '*', '[', ']', ':', '?'}
    sanitized = ''.join(char for char in person_name if char not in invalid_chars).strip() or 'Hoja'
    base_name = sanitized[:31]
    candidate = base_name
    suffix = 1

    while candidate in used_names:
        suffix_text = f"_{suffix}"
        candidate = f"{base_name[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1

    used_names.add(candidate)
    return candidate


def generar_excel(reportes_por_persona, start_date, end_date):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_names = set()
    header_fill = PatternFill(fill_type='solid', fgColor='097F8E')
    header_font = Font(color='FFFFFF', bold=True)
    saldo_font = Font(bold=True)

    for report in reportes_por_persona:
        sheet = workbook.create_sheet(_build_sheet_name(report['name'], used_names))
        sheet.append([f"Caja {report['name']}"])
        sheet.append([f"Desde: {start_date.strftime('%d/%m/%Y')}"])
        sheet.append([f"Hasta: {end_date.strftime('%d/%m/%Y')}"])
        sheet.append([])
        sheet.append(["Fecha", "Concepto", "Descripción", "Cantidad", "Total acumulado"])

        for cell in sheet[5]:
            cell.fill = header_fill
            cell.font = header_font

        sheet.append(["Saldo", "", "", report['saldo_total'], ""])
        for cell in sheet[6]:
            cell.font = saldo_font

        for row in report['rows']:
            sheet.append([
                row['date'],
                row['concept'],
                row['description'],
                row['amount'],
                row['running_total'],
            ])

        for row in sheet.iter_rows(min_row=6, min_col=4, max_col=5):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00 [$EUR]'

        widths = {
            'A': 14,
            'B': 28,
            'C': 40,
            'D': 15,
            'E': 18,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
